"""Kabul testi. Çalıştırma: backend klasöründe `python -m pytest tests` veya `python tests/test_acceptance.py`.

Tüm veriler uydurmadır; TC/IBAN algoritmik olarak geçerli ama kimseye ait değildir.
"""

import sys
import tempfile
from pathlib import Path

sys.path.insert(0, str(Path(__file__).resolve().parents[1]))

from kalkan import llm  # noqa: E402
from kalkan.detectors import iban_valid, luhn_valid, tc_valid, vkn_valid  # noqa: E402
from kalkan.engine import Engine  # noqa: E402

ORNEK = (
    "Müşterimiz Ahmet Yılmaz (TC: 12345678950) iade istiyor. Telefon 0532 123 45 67,\n"
    "e-posta ahmet.yilmaz@gmail.com. İade TR18 0006 2001 1900 0006 6723 15 IBAN'ına.\n"
    "Adresi Cumhuriyet Mahallesi Gazi Cad. No: 12 Elazığ. Ayşe Hanım aradı.\n"
    "Sipariş no 12345678901. Kendisi şeker hastası; kâr marjımız %18."
)


def _engine() -> Engine:
    return Engine(Path(tempfile.mkdtemp()))


def test_validators():
    assert tc_valid("12345678950")
    assert not tc_valid("12345678901")
    assert iban_valid("TR18 0006 2001 1900 0006 6723 15")
    assert not iban_valid("TR18 0006 2001 1900 0006 6723 16")
    assert luhn_valid("4111 1111 1111 1111")
    assert not luhn_valid("4111 1111 1111 1112")
    assert vkn_valid("4840847211")
    assert not vkn_valid("4840847212")


def test_acceptance_regex_layer():
    e = _engine()
    r = e.mask(ORNEK, session="t", use_llm=False)
    masked, by_type = r["masked"], {}
    for f in r["findings"]:
        by_type.setdefault(f["type"], []).append(f["value"])

    assert by_type["TC"] == ["12345678950"]
    assert by_type["IBAN"] == ["TR18 0006 2001 1900 0006 6723 15"]
    assert by_type["TELEFON"] == ["0532 123 45 67"]
    assert by_type["EPOSTA"] == ["ahmet.yilmaz@gmail.com"]
    assert by_type["ADRES"] == ["Cumhuriyet Mahallesi Gazi Cad. No: 12 Elazığ"]
    assert by_type["KISI"] == ["Ahmet Yılmaz", "Ayşe"]
    assert "12345678901" in masked, "sipariş no maskelenmemeli"
    assert "[KISI_2] Hanım" in masked
    for raw in ("Ahmet", "12345678950", "0532", "gmail", "TR18", "Cumhuriyet"):
        assert raw not in masked
    assert e.unmask(masked, "t") == ORNEK


def test_roundtrip_and_consistent_tags():
    e = _engine()
    text = "Ahmet Yılmaz aradı. Ahmet Yılmaz yarın tekrar arayacak, numarası 0532 123 45 67."
    r = e.mask(text, session="s", use_llm=False)
    assert r["masked"].count("[KISI_1]") == 2
    assert e.unmask(r["masked"], "s") == text
    # AI cevabındaki etiketler geri çevrilir.
    assert e.unmask("[KISI_1] için iade onaylandı.", "s") == "Ahmet Yılmaz için iade onaylandı."
    # Başka oturumun etiketleri çözülmez.
    assert e.unmask("[KISI_1]", "baska") == "[KISI_1]"


def test_nonbreaking_spaces():
    # Word/PDF/web'den kopyalanan metinlerde boşluklar U+00A0 olabilir; IBAN yine yakalanmalı.
    e = _engine()
    text = "IBAN: TR18 0006 2001 1900 0006 6723 15, Ahmet Yılmaz"
    r = e.mask(text, session="nb", use_llm=False)
    assert "TR18" not in r["masked"] and "Ahmet" not in r["masked"], r["masked"]
    assert e.unmask(r["masked"], "nb") == text


def test_type_filter_and_policy():
    e = _engine()
    # Kişi adları filtre dışı: isim açık kalır, TC yine gizlenir.
    r = e.mask(ORNEK, session="f1", use_llm=False, types=["TC", "IBAN", "TELEFON", "EPOSTA", "ADRES"])
    assert "Ahmet Yılmaz" in r["masked"] and "12345678950" not in r["masked"]
    assert "KISI" in r["excluded"]
    # Boş liste: hiçbir şey gizlenmez.
    assert e.mask(ORNEK, session="f2", use_llm=False, types=[])["masked"] == ORNEK
    # Filtre çakışma çözümünden önce uygulanır: adres kapalıyken adresin içindeki isim yine gizlenir.
    r = e.mask("Ahmet Yılmaz Sokak No: 5", session="f3", use_llm=False, types=["KISI"])
    assert r["masked"].startswith("[KISI_1]"), r["masked"]
    # Şirket politikası, istekte tür listesi yokken geçerlidir.
    e.policy.set(["TC"])
    r = e.mask(ORNEK, session="f4", use_llm=False)
    assert set(r["counts"]) == {"TC"}, r["counts"]


def test_file_review_and_selection():
    # İki adımlı akış: önce incele (dosya değişmez), sonra yalnızca seçilenleri gizle.
    import io as _io
    import docx
    from kalkan.engine import normalize
    from kalkan.files import mask_file

    sample = Path(__file__).resolve().parents[2] / "ornek" / "sozlesme.docx"
    if not sample.exists():
        print("  (ornek/sozlesme.docx yok — python ornek/olustur.py ile üretin; test atlandı)")
        return
    e = _engine()
    found, cache = [], {}
    mask_file(e, "s.docx", sample.read_bytes(), "analiz", False, found=found, phrase_cache=cache, record=False)
    pairs = {(t, v) for t, v in found}
    assert ("TC", "10000000146") in pairs and ("KISI", "Ayşe Demir") in pairs, pairs
    only = {("TC", normalize("TC", "10000000146")), ("KISI", normalize("KISI", "Ayşe Demir"))}
    out, _, rep = mask_file(e, "s.docx", sample.read_bytes(), "secim", False, ["TC", "KISI"], only=only,
                            phrase_cache=cache, record=False)
    d = docx.Document(_io.BytesIO(out))
    text = "\n".join(p.text for p in d.paragraphs) + "\n".join(c.text for t in d.tables for r in t.rows for c in r.cells)
    assert "10000000146" not in text and "Ayşe Demir" not in text  # seçilenler gizli
    assert "Mehmet" in text and "0532 123 45 67" in text  # seçilmeyenler açık
    assert rep["toplam"] == 2, rep


def test_form_fields_and_names_do_not_cross_lines():
    from kalkan.detectors import detect_names
    e = _engine()
    # Resmî belge düzeni (uydurma veri): etiket, ':' ve değer ayrı satırlarda da olabilir.
    split = "Adı\n:\nMUHAMMED SAFA\nSoyadı\n:\nYİĞİT\nBaba Adı\n:\nMAHMUT\nDoğum Tarihi\n:\n23.01.2007\nSosyal Güvenlik Kurumu"
    same = "T.C. Kimlik No : 12345678950\nAdı : MUHAMMED SAFA\nSoyadı : YİĞİT\nDoğum Yeri : BATMAN\nSayı :\nKonu : Tespit\nTarih: 10.10.2026"
    for text in (split, same):
        r = e.mask(text, session="form", use_llm=False)
        got = {(f["type"], f["value"]) for f in r["findings"]}
        assert ("KISI", "MUHAMMED SAFA") in got and ("KISI", "YİĞİT") in got, got
        assert not any("\n" in v for _, v in got), got  # hiçbir değer satır atlamamalı
    assert ("DOGUM_TARIHI", "23.01.2007") in {(f["type"], f["value"]) for f in e.mask(split, session="f", use_llm=False)["findings"]}
    r = e.mask(same, session="g", use_llm=False)
    assert "10.10.2026" in r["masked"] and "Tespit" in r["masked"]  # belge tarihi ve konu gizlenmez
    assert "BATMAN" not in r["masked"]
    # İsim tespiti: "YİĞİT\nBaba" ve "MAHMUT Sosyal" birleştirilmemeli.
    assert all("\n" not in f.value for f in detect_names("Soyadı\nYİĞİT\nBaba Adı\nMAHMUT Sosyal Güvenlik"))
    assert [f.value for f in detect_names("AHMET Yılmaz geldi")] == []  # biçim tutmuyor, tek büyük harfli kelime


def test_official_pdf_sample():
    from kalkan.files import mask_file
    import pymupdf
    sample = Path(__file__).resolve().parents[2] / "ornek" / "resmi_belge.pdf"
    if not sample.exists():
        print("  (ornek/resmi_belge.pdf yok — python ornek/olustur.py ile üretin; test atlandı)")
        return
    e = _engine()
    out, _, rep = mask_file(e, "r.pdf", sample.read_bytes(), "pdf", False, record=False)
    text = "\n".join(p.get_text() for p in pymupdf.open(stream=out, filetype="pdf"))
    for secret in ("12345678950", "MUHAMMED", "YİĞİT", "MAHMUT", "BATMAN", "23.01.2007"):
        assert secret not in text, secret
    assert "10.10.2026" in text and rep["kalan"] == 0


def test_policy_new_types_default_on():
    import json, tempfile as _t
    d = Path(_t.mkdtemp())
    (d / "politika.json").write_text(json.dumps({"turler": ["TC", "KISI"]}), encoding="utf-8")  # eski biçim
    e = Engine(d)
    assert "DOGUM_TARIHI" in e.policy.types and "IBAN" not in e.policy.types


def test_developer_secrets():
    # Sahte anahtarlar çalışma anında birleştirilir; repoda hazır bir anahtar dizgisi durmaz.
    j = "".join
    keys = {
        "aws": j(["AKIA", "Q3ZXR7VBN2", "KD5TYU"]),
        "github": j(["gh", "p_", "a1B2c3D4e5F6g7H8i9J0k1L2m3N4o5P6q7R8"]),
        "openai": j(["sk-", "proj-", "Xq8Lr2mN7vT4kPz9aB3cD6eF1gH5jK0w"]),
        "anthropic": j(["sk-", "ant-", "api03-", "Zr7Tq2Lm9Kx4Vb8Nc1Wd6Ys3"]),
        "google": j(["AIza", "SyB1c2D3e4F5g6H7i8J9k0L1m2N3o4P5q6R"]),
        "stripe": j(["sk_", "live_", "51Hx2Kq9Lm3Np8Rt"]),
        "jwt": j(["eyJhbGciOiJIUzI1NiJ9", ".", "eyJzdWIiOiIxMjM0NTYiLCJyb2xlIjoiYWRtaW4ifQ", ".", "s1gn4tur3Xyz0123456789abcd"]),
    }
    pem = j(["-----BEGIN ", "RSA PRIVATE KEY-----\nMIIEowIBAAKCAQEA7x9Qk2v\nb3Jt12345678950Zq\n", "-----END RSA ", "PRIVATE KEY-----"])
    text = (
        f"AWS_ACCESS_KEY_ID={keys['aws']}\n"
        f"GITHUB_TOKEN={keys['github']}\n"
        f"OPENAI_API_KEY={keys['openai']}\n"
        f"anthropic = \"{keys['anthropic']}\"\n"
        f"maps key: {keys['google']}\n"
        f"stripe.api_key = '{keys['stripe']}'\n"
        f"Authorization: Bearer {keys['jwt']}\n"
        "DATABASE_URL=postgres://app_user:S3cr3tPa55@10.0.4.12:5432/musteri_db\n"
        "Server=sql01;Database=crm;User Id=sa;Password=Gizli!2026;\n"
        "smtp_password = \"Kalkan2026!\"\n"
        "Yöneticinin şifresi: Elazig.2026\n"
        f"{pem}\n"
        "api_key = <API_KEY>\napi_key = ${API_KEY}\npassword: ********\npassword=changeme\n"
    )
    e = _engine()
    r = e.mask(text, session="sec", use_llm=False)
    masked = r["masked"]
    for name, k in keys.items():
        assert k not in masked, f"{name} anahtarı açıkta kaldı"
    for raw in ("S3cr3tPa55", "Gizli!2026", "Kalkan2026!", "Elazig.2026", "MIIEowIBAAK", "12345678950"):
        assert raw not in masked, raw
    # Anahtar adı bağlam için kalır, yalnızca değer gizlenir.
    assert "smtp_password = \"[SIFRE_" in masked and "şifresi: [SIFRE_" in masked, masked
    # Yer tutucular gizlenmez.
    for placeholder in ("<API_KEY>", "${API_KEY}", "********", "password=changeme"):
        assert placeholder in masked, placeholder
    # PEM içindeki 11 rakam TC sanılıp anahtarı bölmez: blok tek parça gizlenir.
    assert "[TC_" not in masked and masked.count("[GIZLI_ANAHTAR_") == 8, masked
    assert {"BAGLANTI", "SIFRE", "GIZLI_ANAHTAR"} <= set(r["counts"]), r["counts"]
    assert e.unmask(masked, "sec") == text


def test_image_grounding_does_not_overreach():
    # Gemma'dan gelen bir numara OCR kelimelerine bağlanırken başlık satırını ya da komşu sayıları yutmamalı.
    from kalkan.images import Word, ground_phrases
    rows = [["Ad", "Soyad", "TC", "Kimlik", "Telefon"], ["Ahmet", "Yılmaz", "12345678950", "0532", "123", "45", "67"]]
    words, text, pos = [], "", 0
    for li, row in enumerate(rows):
        for wi, t in enumerate(row):
            if wi:
                text += " "
                pos += 1
            words.append(Word(0, 0, 1, 1, t, pos, pos + len(t), li))
            text += t
            pos += len(t)
        text += "\n"
        pos += 1
    found = ground_phrases([{"tur": "TC", "ifade": "12345678950"}, {"tur": "TELEFON", "ifade": "0532 123 45 67"},
                            {"tur": "KISI", "ifade": "Ahmet Yilmaz"}], words, text)
    got = {f.type: f.value for f in found}
    assert got == {"TC": "12345678950", "TELEFON": "0532 123 45 67", "KISI": "Ahmet Yılmaz"}, got


def test_screenshot_redaction():
    from kalkan.files import mask_file
    try:
        from kalkan.images import _tesseract
        _tesseract()
    except Exception:
        print("  (Tesseract yok — ekran görüntüsü testi atlandı)")
        return
    sample = Path(__file__).resolve().parents[2] / "ornek" / "ekran_tablo.png"
    if not sample.exists():
        print("  (ornek/ekran_tablo.png yok — python ornek/olustur.py ile üretin; test atlandı)")
        return
    e = _engine()
    found = []
    _, media, rep = mask_file(e, "ekran.png", sample.read_bytes(), "img", False, found=found, record=False)
    assert media == "image/png" and rep["kalan"] == 0, rep
    assert rep["bulgular"] == {"KISI": 3, "TC": 3, "TELEFON": 3, "IBAN": 3}, rep["bulgular"]
    assert ("IBAN", "TR85 2311 4077 4654 4541 4752 93") in found  # son grup dahil, tam sınır


def test_overlap_loser_is_not_dropped():
    # Model uzun bir SENDIKA ifadesi verir, isim katmanı içinden "Kerem Aksoy"u bulur: isim KISI olur,
    # cümlenin geri kalanı yine SENDIKA olarak gizlenir; "Sendika üyesi" bilgisi açıkta kalmaz.
    from kalkan.detectors import Finding
    e = _engine()
    text = "Sendika üyesi olduğu için yöneticisi Kerem Aksoy çekinceli."
    long = "Sendika üyesi olduğu için yöneticisi Kerem Aksoy çekinceli"
    r = e.mask(text, session="ov", use_llm=False, extra=[Finding(0, len(long), "SENDIKA", long, "llm")])
    assert "Sendika" not in r["masked"] and "Kerem" not in r["masked"], r["masked"]
    assert r["counts"] == {"KISI": 1, "SENDIKA": 2}, r["counts"]
    assert e.unmask(r["masked"], "ov") == text


def test_stats_have_no_raw_data():
    e = _engine()
    e.mask(ORNEK, session="x", use_llm=False)
    raw = e.stats.path.read_text(encoding="utf-8")
    for secret in ("Ahmet", "12345678950", "TR18", "gmail"):
        assert secret not in raw


def test_acceptance_gemma_layer():
    if not llm.available():
        print("  (Ollama/Gemma kapalı — Gemma kabul testi atlandı)")
        return
    e = _engine()
    r = e.mask(ORNEK, session="g", use_llm=True)
    types = {f["value"]: f["type"] for f in r["findings"]}
    assert r["llm_used"]
    assert any(t == "SAGLIK" and "şeker hastası" in v for v, t in types.items()), types
    assert any(t == "TICARI_SIR" and "%18" in v for v, t in types.items()), types
    assert "12345678901" in r["masked"]
    assert e.unmask(r["masked"], "g") == ORNEK


if __name__ == "__main__":
    for name, fn in list(globals().items()):
        if name.startswith("test_") and callable(fn):
            fn()
            print(f"OK  {name}")
    print("Tüm kabul testleri geçti.")
