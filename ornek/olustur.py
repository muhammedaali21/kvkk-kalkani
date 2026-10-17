"""Dosya gizleme denemesi için örnek dosyalar üretir:
ornek/musteriler.xlsx, .csv (müşteri listesi), ornek/sozlesme.docx (Word), ornek/dilekce.pdf (PDF).

Tüm veriler uydurmadır; TC/IBAN algoritmik olarak geçerli ama kimseye ait değildir.
Çalıştırma: python ornek/olustur.py
"""

import random
import sys
from pathlib import Path

import pandas as pd

sys.path.insert(0, str(Path(__file__).resolve().parents[1] / "benchmark"))
import generate as g  # noqa: E402

g.rng.seed(19)
rng = g.rng

KISILER = ["Ahmet Yılmaz", "Ayşe Demir", "Mehmet Kaya", "Zeynep Çelik", "Berkay Kurt",
           "Elif Aydın", "Doruk Aslan", "Merve Şahin"]
NOTLAR = ["", "", "Kargo gecikti, özür e-postası atılacak.", "Kendisi şeker hastası, kapıya teslim istiyor.",
          "Bu müşteriye özel iskonto %30 uygulanıyor.", "İade talebi var.", "Ahmet Bey'in yönlendirmesiyle geldi."]

satirlar = []
# Aynı müşteri birden fazla sipariş verir: maskeli dosyada da aynı etiketi almalı.
for i in range(16):
    kisi = KISILER[i % len(KISILER)] if i < 8 else rng.choice(KISILER)
    rng_k = random.Random(kisi)  # aynı kişi → aynı TC/telefon/IBAN
    g.rng = rng_k
    tc, tel, iban, mail = g.tc(), g.phone(), g.iban(), g.email(kisi)
    g.rng = rng
    satirlar.append({
        "Sipariş No": g.invalid_11(),
        "Tarih": f"{rng.randint(1, 28):02d}.10.2026",
        "Ad Soyad": kisi,
        "TC Kimlik": tc,
        "Telefon": tel,
        "E-posta": mail,
        "IBAN": iban,
        "Şehir": rng.choice(g.SEHIR),
        "Adet": rng.randint(1, 5),
        "Tutar (TL)": rng.randint(150, 9500),
        "Not": rng.choice(NOTLAR),
    })

df = pd.DataFrame(satirlar)
out = Path(__file__).parent
df.to_excel(out / "musteriler.xlsx", index=False)
df.to_csv(out / "musteriler.csv", index=False, sep=";", encoding="utf-8-sig")
print(df.to_string())


# --------------------------------------------------------------------------- Word sözleşme
def word_sozlesme(path: Path) -> None:
    """Bilerek zor durumlar içerir: parçalı (kalın) isim, tablo, e-posta bağlantısı, üst bilgi, yazar bilgisi."""
    import docx
    from docx.opc.constants import RELATIONSHIP_TYPE as RT
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn

    d = docx.Document()
    d.core_properties.author = "Selin Aksoy"
    d.sections[0].header.paragraphs[0].text = "Yıldız Lojistik A.Ş. · Hazırlayan: Selin Aksoy"

    d.add_heading("Hizmet Sözleşmesi", level=1)
    p = d.add_paragraph("İşbu sözleşme, bir tarafta ")
    p.add_run("Mehmet ")
    p.add_run("Kaya").bold = True  # isim iki ayrı biçim parçasına bölünmüş
    p.add_run(" (T.C. 10000000146), diğer tarafta Yıldız Lojistik A.Ş. (vergi kimlik no 4840847211) arasında imzalanmıştır.")
    d.add_paragraph("Müşterinin adresi Atatürk Mahallesi Lale Sokak No: 7 Merkez/Elazığ olup, "
                    "ödemeler TR18 0006 2001 1900 0006 6723 15 IBAN numaralı hesaba yapılacaktır.")
    d.add_paragraph("Müşteri, kalp ameliyatı geçirdiğini ve ağır yük taşıyamayacağını beyan etmiştir. "
                    "Bu müşteriye özel iskonto %30 uygulanacaktır.")

    p = d.add_paragraph("Yazışmalar için: ")
    rid = p.part.relate_to("mailto:mehmet.kaya@ornekfirma.com.tr", RT.HYPERLINK, is_external=True)
    link = OxmlElement("w:hyperlink")
    link.set(qn("r:id"), rid)
    run = OxmlElement("w:r")
    t = OxmlElement("w:t")
    t.text = "mehmet.kaya@ornekfirma.com.tr"
    run.append(t)
    link.append(run)
    p._p.append(link)

    d.add_heading("Taraf bilgileri", level=2)
    table = d.add_table(rows=3, cols=2)
    table.style = "Table Grid"
    for row, (k, v) in zip(table.rows, [("Yetkili", "Ayşe Demir"), ("Telefon", "0532 123 45 67"),
                                        ("Sipariş no", "12345678901")]):
        row.cells[0].text, row.cells[1].text = k, v
    d.add_paragraph("Sözleşme süresi 12 aydır. Teslimat 17.10.2026 tarihinde Elazığ deposundan yapılacaktır.")
    d.save(path)


# --------------------------------------------------------------------------- PDF dilekçe
def pdf_dilekce(path: Path) -> None:
    """Adres iki satıra bölünmüş; karartmanın bölünmüş değerleri de bulması gerekir."""
    import pymupdf

    font = Path("C:/Windows/Fonts/arial.ttf")
    doc = pymupdf.open()
    page = doc.new_page()
    lines = [
        "ELAZIĞ 2. ASLİYE HUKUK MAHKEMESİ'NE",
        "",
        "Davacı: Zeynep Çelik (T.C. 12345678950)",
        "Vekili: Av. Kerem Aksoy, telefon 0544 987 65 43, e-posta kerem.aksoy@hukukburosu.com.tr",
        "Davalı: Yıldız Lojistik A.Ş.",
        "",
        "Müvekkilin adresi Cumhuriyet Mahallesi Gazi Cad. No: 12",
        "Elazığ olup, tazminat bedeli TR18 0006 2001 1900 0006 6723 15 IBAN'ına ödenmelidir.",
        "Müvekkil olay nedeniyle depresyon tedavisi görmektedir.",
        "",
        "Dosya no 2026/1453 olup duruşma 17.10.2026 tarihindedir.",
    ]
    y = 72
    for line in lines:
        page.insert_text((60, y), line, fontsize=11, fontname="arial", fontfile=str(font))
        y += 20
    doc.set_metadata({"author": "Kerem Aksoy", "title": "Dilekçe"})
    doc.subset_fonts()
    doc.save(path, garbage=3, deflate=True)


# --------------------------------------------------------------------------- Resmî belge (form düzeni)
def pdf_resmi_belge(path: Path) -> None:
    """e-Devlet / SGK yazılarına benzer form düzeni: etiket, ':' ve değer PDF'te ayrı ayrı yazılır.
    Kurum adı uydurmadır ve kişi bilgileri tamamen sahtedir; gerçek bir belgenin kopyası değildir."""
    import pymupdf

    reg, bold = "C:/Windows/Fonts/times.ttf", "C:/Windows/Fonts/timesbd.ttf"
    doc = pymupdf.open()
    page = doc.new_page()

    def put(x, y, s, b=False, size=11):
        page.insert_text((x, y), s, fontsize=size, fontname="tb" if b else "tr", fontfile=bold if b else reg)

    for i, line in enumerate(["T.C.", "ÖRNEK KAMU KURUMU BAŞKANLIĞI", "ÖRNEK HİZMETLER GENEL MÜDÜRLÜĞÜ"]):
        put(170 + (60 if i == 0 else 0), 50 + i * 16, line, True, 12)
    put(470, 120, "10.10.2026")
    put(40, 150, "Sayı", True); put(140, 150, ":")
    put(40, 166, "Konu", True); put(140, 166, ":"); put(152, 166, "Sigortalılık Tescil Kaydı Tespiti")
    put(240, 196, "İLGİLİ MAKAMA")
    rows = [("T.C. Kimlik No", "12345678950"), ("Adı", "MUHAMMED SAFA"), ("Soyadı", "YİĞİT"),
            ("Baba Adı", "MAHMUT"), ("Doğum Yeri", "BATMAN"), ("Doğum Tarihi", "23.01.2007")]
    for k, (label, value) in enumerate(rows):
        y = 226 + k * 16
        put(40, y, label, True); put(140, y, ":"); put(152, y, value)
    put(60, 340, "Kurum kayıtlarının incelenmesinde, yukarıda kimlik bilgileri bulunan şahsın 5510 sayılı")
    put(40, 356, "Kanunun 4 üncü maddesine tabi sigortalılığının bulunmadığı tespit edilmiştir.")
    put(60, 372, "Bilgi edinilmesini arz ederim.")
    put(40, 700, "Not: Bu belge sorgulama zamanı olan 10.10.2026 19:34:42 itibariyle düzenlenmiştir.", size=8)
    doc.subset_fonts()
    doc.save(path, garbage=3, deflate=True)


word_sozlesme(out / "sozlesme.docx")
pdf_dilekce(out / "dilekce.pdf")
pdf_resmi_belge(out / "resmi_belge.pdf")
print("sozlesme.docx, dilekce.pdf ve resmi_belge.pdf oluşturuldu.")
